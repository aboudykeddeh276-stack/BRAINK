from pathlib import Path
from openpyxl import load_workbook

class WorkbookService:
    def read_rows(self,path:str,sheet:str,min_row:int=1,max_row:int|None=None):
        wb=load_workbook(path,data_only=False,read_only=True); ws=wb[sheet]
        return [list(r) for r in ws.iter_rows(min_row=min_row,max_row=max_row,values_only=True)]
    def append_row(self,path:str,sheet:str,values:list):
        p=Path(path); wb=load_workbook(p); ws=wb[sheet]; ws.append(values); wb.save(p)
        return {'path':str(p),'sheet':sheet,'row':ws.max_row}
